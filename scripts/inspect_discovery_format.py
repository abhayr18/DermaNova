import docx

doc = docx.Document('d:/DermaNova/reports/1-17 A Review on the Utilization of Cow Dung in Sustainable and Eco-Friendly Building Materials - DONE.docx')
for i in [0, 2, 4, 5, 7, 9, 10, 21, 22]:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None
        sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None
        ls = p.paragraph_format.line_spacing
        print(f'P[{i}]: text="{p.text[:35]}" sb={sb} sa={sa} ls={ls} align={p.alignment}')
